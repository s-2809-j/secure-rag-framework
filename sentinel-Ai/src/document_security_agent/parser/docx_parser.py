from __future__ import annotations

from pathlib import Path

from docx import Document

from src.document_security_agent.exceptions import (
    DocumentParseError,
)
from src.document_security_agent.models import (
    DocumentParseResult,
    FileMetadata,
    ParsedDocument,
)

from .base_parser import BaseParser


class DocxParser(BaseParser):
    """
    Parser for DOCX documents.
    """

    def parse(
        self,
        file_path: Path,
    ) -> DocumentParseResult:

        try:
            doc = Document(file_path)

            text = "\n".join(
                paragraph.text
                for paragraph in doc.paragraphs
            )

            metadata = FileMetadata(
                filename=file_path.name,
                extension=".docx",
                mime_type=(
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml.document"
                ),
                size_bytes=file_path.stat().st_size,
            )

            document = ParsedDocument(
                metadata=metadata,
                text=text,
                page_count=1,
                parser_name=self.__class__.__name__,
            )

            return DocumentParseResult(
                success=True,
                document=document,
            )

        except Exception as exc:
            raise DocumentParseError(
                f"Failed to parse DOCX: {file_path.name}"
            ) from exc